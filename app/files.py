import csv
import io
import re
from itertools import islice
from pathlib import Path


def _table_requirements(rows):
    rows = [[str(c).strip() if c is not None else '' for c in islice(row, 30)] for row in islice(rows, 201)]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    header = [cell.casefold() for cell in rows[0]]
    quantity_keys = ('количество', 'кол-во', 'кол во', 'qty', 'quantity', 'шт', 'саны', 'мөлшері')
    quantity_index = next((i for i, cell in enumerate(header) if any(key in cell for key in quantity_keys)), None)
    name_keys = ('наименование', 'товар', 'название', 'description', 'позиция', 'артикул', 'атауы', 'тауар')
    has_header = quantity_index is not None or any(any(key in cell for key in name_keys) for cell in header)
    output = []
    for row in rows[1:] if has_header else rows:
        qty = 1
        index = quantity_index
        if index is None and len(row) > 1 and re.fullmatch(r'\d+(?:\.0)?', row[-1]):
            index = len(row) - 1
        if index is not None and index < len(row) and row[index]:
            if not re.fullmatch(r'\d+(?:\.0)?', row[index]):
                raise ValueError('Quantity must be a positive integer')
            qty = int(float(row[index]))
            if not 1 <= qty <= 10000:
                raise ValueError('Quantity must be between 1 and 10000')
        description = ' '.join(cell for i, cell in enumerate(row) if cell and i != index)[:400]
        if len(description) >= 3:
            output.append({'description': description, 'quantity': qty})
    return output[:100]


def _text_requirements(lines):
    output = []
    for line in islice(lines, 100):
        line = line.strip()
        if len(line) < 3:
            continue
        match = re.search(r'(?:кол[-. ]?во|qty|quantity|саны|шт)\s*[:=]?\s*(\d+)', line, re.I)
        suffix = re.search(r'\b(\d+)\s*(?:шт\.?|дана)\b', line, re.I)
        quantity = int((match or suffix).group(1)) if match or suffix else 1
        if not 1 <= quantity <= 10000:
            raise ValueError('Quantity must be between 1 and 10000')
        output.append({'description': line[:400], 'quantity': quantity})
    return output


def parse_file(filename: str, data: bytes) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext == '.csv':
        text = data.decode('utf-8-sig', errors='replace')
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=',;\t')
            rows = csv.reader(io.StringIO(text), dialect)
        except csv.Error:
            rows = csv.reader(io.StringIO(text))
        return _table_requirements(rows)
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return _table_requirements(book.active.iter_rows(max_row=201, max_col=30, values_only=True))
        finally:
            book.close()
    elif ext == '.xls':
        import xlrd
        sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        return _table_requirements(sheet.row_values(i) for i in range(sheet.nrows))
    elif ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        output = _text_requirements(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables[:20]:
            output += _table_requirements(([c.text for c in row.cells[:30]] for row in table.rows[:201]))
        return output[:100]
    elif ext == '.pdf':
        import pymupdf
        doc = pymupdf.open(stream=data, filetype='pdf')
        if doc.page_count > 30 or doc.is_encrypted:
            doc.close()
            raise ValueError('PDF exceeds 30 pages or is encrypted')
        lines = [line.strip()[:400] for page in doc for line in page.get_text().splitlines()[:100] if line.strip()]
        doc.close()
        if not lines:
            raise ValueError('PDF has no text layer; scanned PDF requires a configured vision model.')
    elif ext in {'.jpg', '.jpeg', '.png'}:
        raise ValueError('Image recognition requires a configured vision model.')
    else:
        raise ValueError('Unsupported file type')
    return _text_requirements(lines)
